"""
BAKU_MASTER — Generador de archivos.

Genera archivos en formatos estándar de trabajo:
  • DOCX  — Word (propuestas, reportes, briefs)
  • XLSX  — Excel (modelos financieros, planillas de leads, métricas)
  • PDF   — PDF (documentos para clientes, reportes ejecutivos)
  • TXT   — Texto plano
  • MD    — Markdown

Los archivos se guardan en generated_files/ y quedan disponibles para
descarga en el endpoint GET /files/{filename}.
"""

import re
import uuid
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

TZ = ZoneInfo("America/Santiago")
OUTPUT_DIR = Path(__file__).parent.parent / "generated_files"
OUTPUT_DIR.mkdir(exist_ok=True)


def _slug(title: str) -> str:
    """Convierte título a nombre de archivo seguro."""
    s = re.sub(r"[^\w\s-]", "", title.lower())
    s = re.sub(r"[\s_-]+", "_", s).strip("_")
    return s[:50] or "archivo"


def _timestamp() -> str:
    return datetime.now(TZ).strftime("%Y%m%d_%H%M")


def generate_docx(title: str, content: str, subtitle: str = "") -> Path:
    """
    Genera un .docx con estilo profesional.
    Soporta Markdown básico: # encabezados, **negrita**, - listas, ---
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise RuntimeError("python-docx no está instalado. Ejecuta: pip install python-docx")

    doc = Document()

    # Configurar márgenes (2.54cm)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = 914400  # 1 pulgada en EMU

    # Título principal
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.runs[0] if h.runs else h.add_run(title)
    run.font.color.rgb = RGBColor(0x4A, 0x2B, 0xD4)

    # Subtítulo / fecha
    now = datetime.now(TZ)
    date_str = now.strftime("%d/%m/%Y — %H:%M hrs (Santiago)")
    sub_text = f"{subtitle}  |  {date_str}" if subtitle else date_str
    sub = doc.add_paragraph(sub_text)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(10)
    sub.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()  # espacio

    # Procesar contenido línea por línea
    for line in content.split("\n"):
        stripped = line.strip()

        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            p = doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            p = doc.add_heading(stripped[2:], level=1)
        elif stripped == "---" or stripped == "***":
            doc.add_paragraph("─" * 60)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_formatting(p, stripped[2:])
        elif stripped.startswith(tuple(f"{i}. " for i in range(1, 20))):
            p = doc.add_paragraph(style="List Number")
            _add_inline_formatting(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped == "":
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            _add_inline_formatting(p, stripped)

    # Footer
    doc.add_paragraph()
    footer_p = doc.add_paragraph(f"BAKU Agency — Generado el {date_str}")
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.runs[0].font.size = Pt(8)
    footer_p.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    filename = f"{_slug(title)}_{_timestamp()}.docx"
    path = OUTPUT_DIR / filename
    doc.save(str(path))
    return path


def _add_inline_formatting(paragraph, text: str):
    """Procesa **negrita** y *cursiva* dentro de un párrafo."""
    from docx.shared import Pt
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def generate_xlsx(title: str, content: str, sheets: list[dict] | None = None) -> Path:
    """
    Genera un .xlsx con una o más hojas.

    Si `sheets` es None, parsea el contenido como tabla Markdown.
    Si `sheets` es una lista de {name, headers, rows}, crea una hoja por sheet.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl no está instalado. Ejecuta: pip install openpyxl")

    wb = Workbook()
    ws = wb.active

    # Colores BAKU
    PURPLE = "4A2BD4"
    LIGHT_PURPLE = "EDE9FC"
    GRAY = "F5F5F5"
    WHITE = "FFFFFF"

    header_font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    header_fill = PatternFill("solid", fgColor=PURPLE)
    subheader_fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    if sheets:
        ws.title = sheets[0].get("name", "Datos")
        for i, sheet_def in enumerate(sheets):
            if i == 0:
                ws_current = ws
            else:
                ws_current = wb.create_sheet(title=sheet_def.get("name", f"Hoja {i+1}"))
            ws_current.title = sheet_def.get("name", ws_current.title)

            # Título de la hoja
            ws_current.merge_cells(f"A1:{get_column_letter(max(len(sheet_def.get('headers', [])), 1))}1")
            cell = ws_current["A1"]
            cell.value = title
            cell.font = Font(name="Calibri", bold=True, size=14, color=WHITE)
            cell.fill = header_fill
            cell.alignment = center
            ws_current.row_dimensions[1].height = 30

            # Fecha
            now = datetime.now(TZ)
            ws_current["A2"] = f"Generado: {now.strftime('%d/%m/%Y %H:%M')} hrs (Santiago)"
            ws_current["A2"].font = Font(name="Calibri", size=9, color="888888")

            # Headers
            headers = sheet_def.get("headers", [])
            for col_idx, h in enumerate(headers, 1):
                cell = ws_current.cell(row=3, column=col_idx, value=h)
                cell.font = Font(name="Calibri", bold=True, color="333333")
                cell.fill = PatternFill("solid", fgColor="DDDDEE")
                cell.border = border
                cell.alignment = center

            # Filas
            for row_idx, row in enumerate(sheet_def.get("rows", []), 4):
                fill_color = GRAY if row_idx % 2 == 0 else WHITE
                row_fill = PatternFill("solid", fgColor=fill_color)
                for col_idx, val in enumerate(row, 1):
                    cell = ws_current.cell(row=row_idx, column=col_idx, value=val)
                    cell.fill = row_fill
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")

            # Auto-ajustar columnas
            for col in ws_current.columns:
                max_len = 0
                for c in col:
                    if c.value:
                        max_len = max(max_len, len(str(c.value)))
                ws_current.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)
    else:
        # Parsear contenido Markdown como tabla
        ws.title = _slug(title)[:30] or "Datos"
        _parse_markdown_to_xlsx(ws, title, content, header_font, header_fill, border, center, PatternFill, get_column_letter)

    filename = f"{_slug(title)}_{_timestamp()}.xlsx"
    path = OUTPUT_DIR / filename
    wb.save(str(path))
    return path


def _parse_markdown_to_xlsx(ws, title, content, header_font, header_fill, border, center, PatternFill, get_column_letter):
    """Parsea tabla Markdown y la escribe en la hoja."""
    from openpyxl.styles import Font, Alignment, PatternFill as PF
    from openpyxl.utils import get_column_letter as gcl

    # Título
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center

    now = datetime.now(TZ)
    ws["A2"] = f"Generado: {now.strftime('%d/%m/%Y %H:%M')} hrs (Santiago)"
    ws["A2"].font = Font(size=9, color="888888")

    row_num = 3
    in_table = False
    headers_done = False

    for line in content.split("\n"):
        stripped = line.strip()
        if "|" in stripped:
            cols = [c.strip() for c in stripped.split("|") if c.strip()]
            if not cols:
                continue
            if stripped.replace("|", "").replace("-", "").replace(":", "").strip() == "":
                continue  # línea separadora

            for col_idx, val in enumerate(cols, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = border
                cell.alignment = Alignment(vertical="center")
                if not headers_done:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                else:
                    fill_color = "F5F5F5" if row_num % 2 == 0 else "FFFFFF"
                    cell.fill = PF("solid", fgColor=fill_color)
            if not headers_done:
                headers_done = True
            row_num += 1
        elif stripped and not in_table:
            ws.cell(row=row_num, column=1, value=stripped)
            row_num += 1


def generate_pdf(title: str, content: str, subtitle: str = "") -> Path:
    """Genera un PDF limpio y profesional."""
    try:
        from fpdf import FPDF
    except ImportError:
        raise RuntimeError("fpdf2 no está instalado. Ejecuta: pip install fpdf2")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_margins(25, 25, 25)

    # Fuente (FPDF usa fuentes built-in)
    # Título
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(74, 43, 212)  # BAKU purple
    pdf.multi_cell(0, 12, title, align="C")
    pdf.ln(2)

    # Subtítulo + fecha
    now = datetime.now(TZ)
    date_str = now.strftime("%d/%m/%Y — %H:%M hrs (Santiago de Chile)")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(136, 136, 136)
    sub_line = f"{subtitle}  |  {date_str}" if subtitle else date_str
    pdf.multi_cell(0, 6, sub_line, align="C")
    pdf.ln(6)

    # Línea separadora
    pdf.set_draw_color(74, 43, 212)
    pdf.set_line_width(0.8)
    pdf.line(25, pdf.get_y(), 185, pdf.get_y())
    pdf.ln(8)

    # Contenido
    pdf.set_text_color(30, 30, 30)
    for line in content.split("\n"):
        stripped = line.strip()
        if stripped.startswith("### "):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(74, 43, 212)
            pdf.multi_cell(0, 7, stripped[4:])
            pdf.set_text_color(30, 30, 30)
        elif stripped.startswith("## "):
            pdf.ln(4)
            pdf.set_font("Helvetica", "B", 14)
            pdf.set_text_color(74, 43, 212)
            pdf.multi_cell(0, 8, stripped[3:])
            pdf.set_text_color(30, 30, 30)
            pdf.ln(1)
        elif stripped.startswith("# "):
            pdf.ln(5)
            pdf.set_font("Helvetica", "B", 16)
            pdf.set_text_color(74, 43, 212)
            pdf.multi_cell(0, 10, stripped[2:])
            pdf.set_text_color(30, 30, 30)
            pdf.ln(2)
        elif stripped in ("---", "***"):
            pdf.ln(3)
            pdf.set_draw_color(200, 200, 200)
            pdf.set_line_width(0.3)
            pdf.line(25, pdf.get_y(), 185, pdf.get_y())
            pdf.ln(3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            pdf.set_font("Helvetica", "", 10)
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", stripped[2:])
            pdf.multi_cell(0, 6, f"  • {clean}")
        elif stripped == "":
            pdf.ln(4)
        else:
            pdf.set_font("Helvetica", "", 10)
            # Quitar marcas de bold para PDF
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", stripped)
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean)
            pdf.multi_cell(0, 6, clean)

    # Footer
    pdf.ln(8)
    pdf.set_draw_color(74, 43, 212)
    pdf.set_line_width(0.5)
    pdf.line(25, pdf.get_y(), 185, pdf.get_y())
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(136, 136, 136)
    pdf.multi_cell(0, 5, f"BAKU Agency — {date_str}", align="C")

    filename = f"{_slug(title)}_{_timestamp()}.pdf"
    path = OUTPUT_DIR / filename
    pdf.output(str(path))
    return path


def generate_txt(title: str, content: str) -> Path:
    """Genera un archivo de texto plano."""
    now = datetime.now(TZ)
    header = (
        f"{'='*60}\n"
        f"  {title}\n"
        f"  {now.strftime('%d/%m/%Y — %H:%M hrs (Santiago)')}\n"
        f"  BAKU Agency\n"
        f"{'='*60}\n\n"
    )
    filename = f"{_slug(title)}_{_timestamp()}.txt"
    path = OUTPUT_DIR / filename
    path.write_text(header + content, encoding="utf-8")
    return path


def generate_md(title: str, content: str) -> Path:
    """Genera un archivo Markdown."""
    now = datetime.now(TZ)
    header = (
        f"# {title}\n\n"
        f"> Generado: {now.strftime('%d/%m/%Y — %H:%M hrs (Santiago)')} — BAKU Agency\n\n"
        f"---\n\n"
    )
    filename = f"{_slug(title)}_{_timestamp()}.md"
    path = OUTPUT_DIR / filename
    path.write_text(header + content, encoding="utf-8")
    return path


# ── API pública ──────────────────────────────────────────────────────────────

def generate_file(
    title: str,
    content: str,
    file_format: str,
    subtitle: str = "",
    sheets: list[dict] | None = None,
) -> dict:
    """
    Genera un archivo y retorna info del archivo creado.

    file_format: 'docx' | 'xlsx' | 'pdf' | 'txt' | 'md'
    sheets: solo para xlsx — lista de {name, headers, rows}
    """
    fmt = file_format.lower().strip(".")
    try:
        match fmt:
            case "docx":
                path = generate_docx(title, content, subtitle)
            case "xlsx":
                path = generate_xlsx(title, content, sheets)
            case "pdf":
                path = generate_pdf(title, content, subtitle)
            case "txt":
                path = generate_txt(title, content)
            case "md":
                path = generate_md(title, content)
            case _:
                raise ValueError(f"Formato no soportado: {fmt}. Usa: docx, xlsx, pdf, txt, md")

        size_kb = round(path.stat().st_size / 1024, 1)
        return {
            "ok": True,
            "filename": path.name,
            "path": str(path),
            "format": fmt,
            "size_kb": size_kb,
            "download_url": f"/files/{path.name}",
            "created_at": datetime.now(TZ).isoformat(),
        }
    except Exception as e:
        return {"ok": False, "error": str(e), "format": fmt}


def list_files() -> list[dict]:
    """Lista todos los archivos generados."""
    files = []
    for p in sorted(OUTPUT_DIR.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True):
        if p.is_file() and not p.name.startswith("."):
            files.append({
                "filename": p.name,
                "format": p.suffix.lstrip("."),
                "size_kb": round(p.stat().st_size / 1024, 1),
                "download_url": f"/files/{p.name}",
                "created_at": datetime.fromtimestamp(p.stat().st_mtime, TZ).isoformat(),
            })
    return files


def delete_file(filename: str) -> bool:
    """Elimina un archivo generado. Retorna True si se eliminó."""
    path = OUTPUT_DIR / filename
    if path.exists() and path.parent == OUTPUT_DIR:
        path.unlink()
        return True
    return False
